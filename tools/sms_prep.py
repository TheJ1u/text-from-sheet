#!/usr/bin/env python3
"""sms_prep.py - reference implementation + preflight checker for the Excel-to-SMS Shortcut.

What it does (and does NOT do)
------------------------------
DOES:  read a CSV/XLSX of contacts, normalize and validate phone numbers, enforce consent /
       opt-out / suppression / duplicate / resume rules, block sensitive data, render template
       messages, flag content for manual compliance review, and write queue.csv + rejected.csv.
DOES NOT: send anything. This module has no network, SMS, or subprocess code at all
          (tests/test_sms_prep.py enforces that). Sending happens only in the iPhone Shortcut,
          one human-approved message at a time.

The Shortcut implements the same rules natively (see docs/BUILD_SHORTCUT.md). This script is the
tested source of truth for those rules and a way to lint a spreadsheet before running it.

Usage
-----
  python3 tools/sms_prep.py check contacts.csv --config data/config.json \
      --suppression data/suppression_list.csv --log send_log.csv --campaign fall-2026 --out out/
  python3 tools/sms_prep.py init-log send_log.csv
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

LOG_FIELDS = [
    "timestamp", "session_id", "recipient_key", "phone", "name", "template",
    "campaign", "status", "delivery_confirmed", "note",
]
QUEUE_FIELDS = [
    "recipient_key", "first_name", "last_name", "phone", "template", "category",
    "needs_review", "review_reasons", "message",
]
REJECT_FIELDS = ["row", "name", "phone", "reasons"]

# Latest log status per recipient decides what happens on resume.
TERMINAL_STATUSES = {"sent_action_completed", "skipped", "blocked"}   # never re-queue
AMBIGUOUS_STATUSES = {"approved"}   # approved but no outcome logged: may or may not have sent
VALID_LOG_STATUSES = {
    "queued", "approved", "edited", "sent_action_completed", "skipped", "canceled",
    "failed", "blocked",
}

HEADER_ALIASES = {
    "first name": "first_name", "firstname": "first_name", "first": "first_name",
    "last name": "last_name", "lastname": "last_name", "last": "last_name",
    "phone number": "phone", "phone": "phone", "mobile": "phone", "cell": "phone",
    "mobile number": "phone", "cell phone": "phone",
    "message": "message", "notes": "notes", "note": "notes", "status": "status",
    "consent": "consent", "consent date": "consent_date", "consent source": "consent_source",
    "opt out": "opt_out", "optout": "opt_out", "opted out": "opt_out",
    "message type": "message_type", "template": "template",
    "appointment date": "appointment_date", "appointment time": "appointment_time",
    "event name": "event_name", "event date": "event_date",
}

AFFIRMATIVE_CONSENT = {"yes", "y", "true", "1", "opted in", "opt in", "opt-in", "granted"}
OPT_OUT_TRUE = {"yes", "y", "true", "1", "stop", "opted out", "optout", "unsubscribed", "dnc"}
STATUS_SUPPRESSED = {"opted out", "opt out", "do not contact", "dnc", "stop", "unsubscribed", "suppressed"}
STATUS_DONE = {"sent", "done", "complete", "completed", "skipped"}

# --- sensitive data that must never go into an ordinary SMS ---------------------------------
_L = r"[AC-HJKMNP-RT-Y]"          # MBI letters (no S, L, O, I, B, Z)
_LD = r"[AC-HJKMNP-RT-Y0-9]"
SENSITIVE_PATTERNS = {
    "sensitive_ssn": re.compile(r"\b\d{3}-?\d{2}-?\d{4}\b"),
    "sensitive_medicare_id": re.compile(
        rf"\b[1-9]{_L}{_LD}\d-?{_L}{_LD}\d-?{_L}{_L}\d\d\b", re.I),
    "sensitive_card_or_account_number": re.compile(r"\b(?:\d[ -]?){13,19}\b"),
    "sensitive_health_or_financial_term": re.compile(
        r"\b(date of birth|dob|diagnos\w*|prescription\w*|medication\w*|policy number|"
        r"account number|routing number|ssn|social security|bank account|password|pin code)\b",
        re.I),
}

# --- content that is allowed but must be looked at by a human before sending ------------------
RISKY_CLAIM_TERMS = [
    "guaranteed", "free", "no cost", "$0", "save", "savings", "lowest", "best", "cheaper",
    "limited time", "act now", "urgent", "last chance", "you qualify", "you're eligible",
    "you are eligible", "benefits", "$",
]
MEDICARE_TERMS = ["medicare", "medigap", "part d", "part c", "advantage", "enroll", "enrollment", "aep"]


@dataclass
class Result:
    queue: list = field(default_factory=list)
    rejected: list = field(default_factory=list)
    needs_check: list = field(default_factory=list)   # approved-but-unlogged from a prior session
    skipped_done: list = field(default_factory=list)  # already sent/skipped in a prior session
    warnings: list = field(default_factory=list)

    def summary(self) -> dict:
        reasons: dict = {}
        for r in self.rejected:
            for reason in r["reasons"].split(";"):
                reasons[reason] = reasons.get(reason, 0) + 1
        return {
            "ready_for_review": len(self.queue),
            "needs_manual_compliance_review": sum(1 for q in self.queue if q["needs_review"] == "yes"),
            "rejected": len(self.rejected),
            "already_processed_in_prior_session": len(self.skipped_done),
            "ambiguous_prior_approvals": len(self.needs_check),
            "reject_reasons": reasons,
        }


# ------------------------------------------------------------------------------ phone numbers
def normalize_phone(raw: str, default_cc: str = "1"):
    """Return (e164, error). Never guesses a country other than the default NANP '+1'."""
    s = (raw or "").strip()
    if not s:
        return None, "missing_phone"
    if re.search(r"\d[eE][+-]?\d", s):
        return None, "phone_scientific_notation"
    if re.search(r"(ext\.?|x)\s*\d+\s*$", s, re.I):
        return None, "phone_has_extension"
    if re.search(r"[A-Za-z]", s):
        return None, "phone_has_letters"
    has_plus = s.startswith("+") or s.startswith("00")
    digits = re.sub(r"\D", "", s)
    if s.startswith("00"):
        digits = digits[2:]
    if has_plus:
        if not 8 <= len(digits) <= 15:
            return None, "phone_bad_length"
        e164 = "+" + digits
    else:
        if default_cc == "1" and len(digits) == 10:
            e164 = "+1" + digits
        elif default_cc == "1" and len(digits) == 11 and digits.startswith("1"):
            e164 = "+" + digits
        else:
            return None, "phone_needs_country_code"
    if e164.startswith("+1"):
        if not re.fullmatch(r"\+1[2-9]\d{2}[2-9]\d{6}", e164):
            return None, "phone_invalid_nanp"
    return e164, None


# ------------------------------------------------------------------------------ file reading
def canon_header(h: str) -> str:
    c = re.sub(r"[\s_\-]+", " ", (h or "").strip().lower())
    return HEADER_ALIASES.get(c, c)


def read_contacts(path: Path):
    """Return list of (row_number, dict-of-canonical-columns, extra_columns_flag)."""
    path = Path(path)
    rows = []
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            import openpyxl  # optional dependency; CSV path needs nothing
        except ImportError as exc:  # pragma: no cover
            raise SystemExit("Reading .xlsx needs openpyxl (pip install openpyxl). "
                             "Or export the sheet to CSV.") from exc
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        headers = [canon_header(str(h or "")) for h in next(it, [])]
        for i, vals in enumerate(it, start=2):
            d = {}
            for h, v in zip(headers, vals):
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                d[h] = "" if v is None else str(v)
            rows.append((i, d, False))
        return rows
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        try:
            headers = [canon_header(h) for h in next(reader)]
        except StopIteration:
            return rows
        for i, vals in enumerate(reader, start=2):
            if not any(v.strip() for v in vals):
                continue
            extra = len(vals) > len(headers)
            d = {h: (vals[j].strip() if j < len(vals) else "") for j, h in enumerate(headers)}
            rows.append((i, d, extra))
    return rows


def load_suppression(path, default_cc="1") -> set:
    out = set()
    if not path or not Path(path).exists():
        return out
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    if not rows:
        return out
    header = [canon_header(h) for h in rows[0]]
    col = header.index("phone") if "phone" in header else 0
    body = rows[1:] if "phone" in header else rows
    for r in body:
        if len(r) > col and r[col].strip() and not r[col].strip().startswith("#"):
            e164, err = normalize_phone(r[col], default_cc)
            if e164:
                out.add(e164)
    return out


def load_log_state(path) -> dict:
    """Latest status per recipient_key (file order = time order, because the Shortcut only appends)."""
    state = {}
    if not path or not Path(path).exists():
        return state
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            key = (row.get("recipient_key") or "").strip()
            if key:
                state[key] = (row.get("status") or "").strip().lower()
    return state


# ------------------------------------------------------------------------------ checks
def find_sensitive(text: str) -> list:
    return [code for code, pat in SENSITIVE_PATTERNS.items() if pat.search(text or "")]


def review_reasons_for(text: str, category: str) -> list:
    reasons = []
    low = (text or "").lower()
    if category == "marketing":
        reasons.append("marketing_content")
    if any(t in low for t in RISKY_CLAIM_TERMS):
        reasons.append("claims_or_promotional_language")
    if any(re.search(rf"\b{re.escape(t)}\b", low) for t in MEDICARE_TERMS):
        reasons.append("medicare_reference_cms_rules")
    if len(text or "") > 320:
        reasons.append("long_message")
    return reasons


def parse_date(s: str):
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(s.strip(), fmt).date()
        except ValueError:
            continue
    return None


def recipient_key(campaign: str, template: str, phone: str) -> str:
    return f"{campaign}|{template}|{phone}"


def render(template: dict, values: dict):
    """Fill placeholders. Returns (text, missing_fields). Never invents a value."""
    missing = [r for r in template["required"] if not (values.get(r) or "").strip()]
    if missing:
        return None, missing
    text = template["body"].format_map({k: (v or "").strip() for k, v in values.items()})
    return re.sub(r"\s+", " ", text).strip(), []


def process(rows, config, templates, suppression, log_state, campaign,
            default_template=None, today=None) -> Result:
    today = today or date.today()
    cc = config.get("default_country_code", "1")
    optout = config.get("optout_text", "")
    res = Result()
    seen_phones = set()

    for row_no, r, extra in rows:
        name = f"{r.get('first_name', '')} {r.get('last_name', '')}".strip()
        phone_raw = r.get("phone", "")
        reasons = []

        def reject(*codes):
            res.rejected.append({"row": row_no, "name": name, "phone": phone_raw,
                                 "reasons": ";".join(codes)})

        if extra:
            reject("malformed_row_extra_columns_check_commas")
            continue

        e164, perr = normalize_phone(phone_raw, cc)
        if perr:
            reasons.append(perr)
        if not r.get("first_name", "").strip():
            reasons.append("missing_first_name")

        # --- consent / opt-out / suppression (never assume eligibility) ---
        if (r.get("consent", "").strip().lower()) not in AFFIRMATIVE_CONSENT:
            reasons.append("no_affirmative_consent")
        if r.get("opt_out", "").strip().lower() in OPT_OUT_TRUE:
            reasons.append("opted_out")
        status = r.get("status", "").strip().lower()
        if status in STATUS_SUPPRESSED:
            reasons.append("status_suppressed")
        if e164 and e164 in suppression:
            reasons.append("on_suppression_list")
        if status in STATUS_DONE:
            reasons.append("status_already_processed")

        # --- template / message resolution ---
        msg_col = r.get("message", "").strip()
        tpl_id = r.get("template", "").strip().lower() or None
        custom_as_is = False
        if msg_col:
            if tpl_id in (None, "custom"):
                custom_as_is = True
                tpl_id = "custom"
            elif tpl_id != "general":
                reasons.append("message_conflicts_with_template")
        elif tpl_id is None:
            tpl_id = default_template
        if tpl_id is None:
            reasons.append("no_template_or_message")
        elif tpl_id != "custom" and tpl_id not in templates:
            reasons.append("unknown_template")

        # --- category: marketing unless explicitly service AND template allows service ---
        mtype = r.get("message_type", "").strip().lower()
        tpl = templates.get(tpl_id) if tpl_id in templates else None
        service_ok = bool(tpl and tpl.get("service_eligible")) or (custom_as_is)
        category = "service" if (mtype == "service" and service_ok) else "marketing"
        if category == "marketing":
            cd = parse_date(r.get("consent_date", "")) if r.get("consent_date") else None
            if cd is None:
                reasons.append("marketing_needs_valid_consent_date")
            elif cd > today:
                reasons.append("consent_date_in_future")
            if not r.get("consent_source", "").strip():
                reasons.append("marketing_needs_consent_source")

        # --- render ---
        message = None
        if not reasons and tpl_id:
            values = {
                "first_name": r.get("first_name", ""), "business_name": config.get("business_name", ""),
                "agent_name": config.get("agent_name", ""), "optout": optout,
                "appointment_date": r.get("appointment_date", ""), "appointment_time": r.get("appointment_time", ""),
                "event_name": r.get("event_name", ""), "event_date": r.get("event_date", ""),
                "custom_message": msg_col,
            }
            if custom_as_is:
                message = re.sub(r"\s+", " ", msg_col).strip()
                if optout and optout.lower() not in message.lower():
                    message = f"{message} {optout}"
            else:
                message, missing = render(tpl, values)
                reasons += [f"missing_field:{m}" for m in missing]
        if message:
            reasons += find_sensitive(message)

        # --- duplicates ---
        if e164 and not reasons:
            if e164 in seen_phones:
                reasons.append("duplicate_phone_in_file")
        if reasons:
            reject(*dict.fromkeys(reasons))
            continue
        seen_phones.add(e164)

        # --- resume against prior sessions ---
        key = recipient_key(campaign, tpl_id, e164)
        prior = log_state.get(key)
        entry = {"recipient_key": key, "first_name": r["first_name"].strip(),
                 "last_name": r.get("last_name", "").strip(), "phone": e164,
                 "template": tpl_id, "category": category}
        if prior in TERMINAL_STATUSES:
            res.skipped_done.append({**entry, "prior_status": prior})
            continue
        if prior in AMBIGUOUS_STATUSES:
            res.needs_check.append({**entry, "prior_status": prior})
            continue

        rr = review_reasons_for(message, category)
        if r.get("notes") and find_sensitive(r["notes"]):
            res.warnings.append(f"row {row_no}: Notes column contains sensitive-looking data; "
                                "Notes are never sent or given to AI, but consider removing it.")
        res.queue.append({**entry, "needs_review": "yes" if rr else "no",
                          "review_reasons": ";".join(rr), "message": message})
    return res


# ------------------------------------------------------------------------------ AI helpers
def ai_minimal_payload(draft_message: str, first_name: str) -> dict:
    """The ONLY data an AI provider should ever receive: a first name and the already-rendered
    draft. No last name, phone number, notes, consent data, or status."""
    return {"first_name": first_name, "draft": draft_message}


AI_SYSTEM_PROMPT = (
    "You lightly polish short text messages for a licensed insurance agent. Use ONLY the facts "
    "already in the draft. Do not add facts, numbers, dates, prices, plan names, benefits, health "
    "or financial details, promises, or urgency. Keep a friendly, professional tone, under 320 "
    "characters. Return only the message text."
)


def ai_output_guard(draft: str, ai_text: str, optout: str = ""):
    """Return (ok, problems). AI output is only ever a draft for human review; this catches
    obvious drift (invented numbers, money, sensitive data, dropped opt-out, too long)."""
    problems = []
    if not ai_text or not ai_text.strip():
        return False, ["empty_output"]
    if len(ai_text) > 320:
        problems.append("too_long")
    if optout and optout.lower() not in ai_text.lower():
        problems.append("opt_out_missing")
    if "$" in ai_text and "$" not in draft:
        problems.append("introduced_money")
    draft_numbers = set(re.findall(r"\d+", draft))
    if set(re.findall(r"\d+", ai_text)) - draft_numbers:
        problems.append("introduced_numbers")
    problems += find_sensitive(ai_text)
    return (not problems), problems


# ------------------------------------------------------------------------------ CLI
def _write_csv(path: Path, fields, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def cmd_check(args) -> int:
    base = Path(__file__).resolve().parent.parent
    cfg_path = Path(args.config) if args.config else base / "data" / "config.example.json"
    config = json.loads(cfg_path.read_text(encoding="utf-8"))
    templates = {k: v for k, v in json.loads(
        Path(args.templates or base / "templates" / "templates.json").read_text(encoding="utf-8")).items()
        if not k.startswith("_")}
    res = process(read_contacts(Path(args.contacts)), config, templates,
                  load_suppression(args.suppression, config.get("default_country_code", "1")),
                  load_log_state(args.log), args.campaign, args.template)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    _write_csv(out / "queue.csv", QUEUE_FIELDS, res.queue)
    _write_csv(out / "rejected.csv", REJECT_FIELDS, res.rejected)
    if res.needs_check:
        _write_csv(out / "needs_manual_check.csv", QUEUE_FIELDS + ["prior_status"], res.needs_check)
    print(json.dumps(res.summary(), indent=2))
    for w in res.warnings:
        print("WARNING:", w, file=sys.stderr)
    print(f"\nWrote {out}/queue.csv and {out}/rejected.csv. Nothing was sent.")
    return 0


def cmd_init_log(args) -> int:
    p = Path(args.path)
    if p.exists():
        print(f"{p} already exists; not overwriting.", file=sys.stderr)
        return 1
    _write_csv(p, LOG_FIELDS, [])
    print(f"Created {p}")
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check", help="validate a contact file and build queue.csv / rejected.csv")
    c.add_argument("contacts")
    c.add_argument("--config")
    c.add_argument("--templates")
    c.add_argument("--suppression")
    c.add_argument("--log", help="existing send_log.csv, to resume without repeating recipients")
    c.add_argument("--campaign", required=True, help="campaign id, e.g. fall-2026-workshop")
    c.add_argument("--template", help="default template id for rows with no Template/Message")
    c.add_argument("--out", default="out")
    c.set_defaults(func=cmd_check)
    i = sub.add_parser("init-log", help="create an empty send_log.csv with the correct header")
    i.add_argument("path")
    i.set_defaults(func=cmd_init_log)
    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
